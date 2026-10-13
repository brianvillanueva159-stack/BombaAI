import io
from pathlib import Path

import streamlit as st
import ollama
import speech_recognition as sr
from pypdf import PdfReader
from docx import Document


# =========================================================
# PAGE CONFIGURATION
# =========================================================

st.set_page_config(
    page_title="BombaAI",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded",
)


# =========================================================
# THEME AND LAYOUT
# =========================================================

st.markdown(
    """
    <style>
    :root {
        color-scheme: dark;
    }

    .stApp {
        background: #101218;
        color: #f4f4f5;
    }

    [data-testid="stHeader"] {
        background: #101218;
    }

    [data-testid="stMainBlockContainer"] {
        max-width: 1000px;
        padding-top: 1rem;
        padding-bottom: 7rem;
    }

    [data-testid="stSidebar"] {
        background: #171a22;
        border-right: 1px solid #292d37;
    }

    [data-testid="stSidebar"] button {
        border-radius: 10px;
    }

    .bomba-title {
        text-align: center;
        font-size: 42px;
        font-weight: 800;
        letter-spacing: -1.5px;
        margin: 8px 0 0 0;
        color: #ffffff;
    }

    .bomba-subtitle {
        text-align: center;
        color: #9ca3af;
        font-size: 14px;
        margin-top: 5px;
        margin-bottom: 28px;
    }

    [data-testid="stChatMessage"] {
        background: #191c24;
        border: 1px solid #292d37;
        border-radius: 16px;
        padding: 14px;
    }

    [data-testid="stChatInput"] {
        border-radius: 22px;
    }

    [data-testid="stChatInput"] > div {
        background: #242832 !important;
        border: 1px solid #393e49 !important;
        border-radius: 22px !important;
        box-shadow: 0 8px 30px rgba(0, 0, 0, 0.22);
    }

    [data-testid="stChatInput"] textarea {
        color: #f4f4f5 !important;
        background: transparent !important;
    }

    [data-testid="stChatInput"] textarea::placeholder {
        color: #9ca3af !important;
    }

    [data-testid="stChatInput"] button {
        border-radius: 50%;
    }

    #MainMenu,
    footer {
        visibility: hidden;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# CONFIGURATION
# =========================================================

SYSTEM_PROMPT = """
You are BombaAI, a helpful AI assistant.

Answer clearly, accurately, and naturally.
Help with programming, schoolwork, writing, explanations,
document analysis, and general questions.

When a document is attached, use its provided text to answer.
When an image is attached, describe what is visible and answer
questions about it. Do not invent details that cannot be seen.
"""

DEFAULT_MODEL = "llama3.2"
VISION_MODEL = "llava"

SUPPORTED_FILE_TYPES = [
    "png", "jpg", "jpeg", "webp", "bmp",
    "pdf", "docx", "txt", "csv", "json", "md",
    "py", "html", "htm", "xml", "yaml", "yml",
    "log", "ini", "toml", "js", "css", "sql",
]

TEXT_EXTENSIONS = {
    ".txt", ".csv", ".json", ".md", ".py",
    ".html", ".htm", ".xml", ".yaml", ".yml",
    ".log", ".ini", ".toml", ".js", ".css", ".sql",
}

IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".webp", ".bmp",
}


# =========================================================
# SESSION STATE
# =========================================================

if "messages" not in st.session_state:
    st.session_state.messages = []

if "model_name" not in st.session_state:
    st.session_state.model_name = DEFAULT_MODEL


# =========================================================
# LEFT SIDEBAR
# =========================================================

with st.sidebar:
    st.markdown("## ⚙️ Settings")
    st.caption("Customize your BombaAI")

    st.divider()

    st.markdown("### AI Model")

    text_models = [
        "llama3.2",
        "llama3.1",
        "mistral",
    ]

    if st.session_state.model_name not in text_models:
        st.session_state.model_name = DEFAULT_MODEL

    st.session_state.model_name = st.selectbox(
        "Choose a text model",
        text_models,
        index=text_models.index(st.session_state.model_name),
    )

    st.caption(
        "Install a model first using ollama pull modelname."
    )

    st.divider()

    if st.button("🗑️ Clear conversation", use_container_width=True):
        st.session_state.messages = []
        st.rerun()

    st.divider()

    st.markdown("### About")
    st.caption("Your personal AI assistant.")
    st.caption("Powered by Python, Streamlit, and Ollama.")


# =========================================================
# FILE READER
# =========================================================

def read_file(uploaded_file):
    """Return extracted text and image bytes, if applicable."""

    filename = uploaded_file.name
    extension = Path(filename).suffix.lower()
    raw_bytes = uploaded_file.getvalue()

    if extension in IMAGE_EXTENSIONS:
        return f"[Image attached: {filename}]", raw_bytes

    if extension == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(raw_bytes))
            pages = []

            for number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""

                if text.strip():
                    pages.append(f"Page {number}:\n{text}")

            extracted = "\n\n".join(pages)

            if not extracted.strip():
                extracted = (
                    "No selectable text was found. "
                    "This may be a scanned PDF that needs OCR."
                )

            return f"Document: {filename}\n\n{extracted}", None

        except Exception as error:
            return f"Could not read PDF {filename}: {error}", None

    if extension == ".docx":
        try:
            document = Document(io.BytesIO(raw_bytes))

            parts = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " | ".join(cell.text.strip() for cell in row.cells)
                    )

            extracted = "\n".join(parts) or "No text found."

            return f"Document: {filename}\n\n{extracted}", None

        except Exception as error:
            return f"Could not read DOCX {filename}: {error}", None

    if extension in TEXT_EXTENSIONS:
        text = raw_bytes.decode("utf-8", errors="replace")
        return f"File: {filename}\n\n{text}", None

    return (
        f"Unsupported file format: {filename}",
        None,
    )


# =========================================================
# VOICE TRANSCRIPTION
# =========================================================

def transcribe_audio(audio_file):
    """
    Transcribe the recorded WAV audio using online recognition.
    An internet connection is required.
    """

    recognizer = sr.Recognizer()

    try:
        with sr.AudioFile(
            io.BytesIO(audio_file.getvalue())
        ) as source:
            audio_data = recognizer.record(source)

        text = recognizer.recognize_google(
            audio_data,
            language="en-US",
        )

        return text, None

    except sr.UnknownValueError:
        return None, "I couldn't understand the recording. Try again."

    except sr.RequestError as error:
        return None, f"Speech recognition connection failed: {error}"

    except Exception as error:
        return None, f"Voice transcription failed: {error}"


# =========================================================
# HEADER
# =========================================================

st.markdown(
    '<div class="bomba-title">🤖 BombaAI</div>',
    unsafe_allow_html=True,
)

st.markdown(
    '<div class="bomba-subtitle">'
    'Your AI assistant · Powered by Ollama'
    '</div>',
    unsafe_allow_html=True,
)


# =========================================================
# DISPLAY CONVERSATION
# =========================================================

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])


# =========================================================
# COMBINED MESSAGE INPUT
# =========================================================

try:
    prompt = st.chat_input(
        "Message BombaAI...",
        key="bomba_chat_input",
        accept_file="multiple",
        accept_audio=True,
        audio_sample_rate=16000,
        file_type=SUPPORTED_FILE_TYPES,
    )

except TypeError:
    st.error(
        "Your Streamlit version is too old. Update it with:\n\n"
        "py -3.12 -m pip install --upgrade streamlit"
    )
    st.stop()


# =========================================================
# PROCESS INPUT
# =========================================================

if prompt is not None:
    user_text = (prompt.text or "").strip()
    uploaded_files = list(prompt.files or [])
    recorded_audio = prompt.audio

    # Convert audio into text.
    if recorded_audio is not None:
        with st.spinner("Transcribing your voice..."):
            recognized_text, voice_error = transcribe_audio(
                recorded_audio
            )

        if voice_error:
            st.warning(voice_error)

        elif recognized_text:
            if user_text:
                user_text += "\n\nVoice command: " + recognized_text
            else:
                user_text = recognized_text

    # Extract text from files and collect image bytes separately.
    document_texts = []
    image_data_list = []
    attachment_names = []

    for uploaded_file in uploaded_files:
        attachment_names.append(uploaded_file.name)

        extracted_text, image_bytes = read_file(uploaded_file)

        if extracted_text:
            document_texts.append(extracted_text)

        if image_bytes is not None:
            image_data_list.append(image_bytes)

    # Create the actual message for Ollama.
    full_user_text = user_text

    if document_texts:
        full_user_text += (
            "\n\nAttached file contents:\n\n"
            + "\n\n".join(document_texts)
        ) if full_user_text else (
            "Please analyze these attached files:\n\n"
            + "\n\n".join(document_texts)
        )

    if image_data_list and not full_user_text.strip():
        full_user_text = (
            "Describe the attached image or images. "
            "Explain visible text and important details."
        )

    # Reject a completely empty submission.
    if not full_user_text.strip():
        st.info("Type a message, attach a file, or record a voice command.")

    else:
        # Choose a vision-capable model whenever images are attached.
        active_model = (
            VISION_MODEL if image_data_list
            else st.session_state.model_name
        )

        display_text = user_text or "Please analyze the attached files."

        if attachment_names:
            display_text += (
                "\n\n📎 Attached: " + ", ".join(attachment_names)
            )

        # Store the conversation text only.
        st.session_state.messages.append({
            "role": "user",
            "content": full_user_text,
        })

        with st.chat_message("user"):
            st.markdown(display_text)

        # Build a fresh API history. Do not modify the stored messages
        # when attaching images to the request.
        history = [
            {
                "role": "system",
                "content": SYSTEM_PROMPT,
            }
        ]

        history.extend(
            {
                "role": message["role"],
                "content": message["content"],
            }
            for message in st.session_state.messages[-10:]
        )

        # Attach images only to the newest user request.
        if image_data_list:
            for message in reversed(history):
                if message["role"] == "user":
                    message["images"] = image_data_list
                    break

        # Call Ollama and display the response.
        with st.chat_message("assistant"):
            try:
                with st.spinner(
                    f"BombaAI is using {active_model}..."
                ):
                    response = ollama.chat(
                        model=active_model,
                        messages=history,
                        stream=True,
                        options={"num_predict": 700},
                    )

                    answer = st.write_stream(
                        chunk["message"]["content"]
                        for chunk in response
                    )

                st.session_state.messages.append({
                    "role": "assistant",
                    "content": answer,
                })

            except Exception as error:
                st.error(
                    "BombaAI could not connect to the selected Ollama "
                    "model. Check that Ollama is running and that the "
                    "model supports the type of input you sent.\n\n"
                    f"Details: {error}"
                )

                # Remove the failed user request from history.
                if (
                    st.session_state.messages
                    and st.session_state.messages[-1]["role"] == "user"
                ):
                    st.session_state.messages.pop()